#!/usr/bin/env python3
"""Offline tests for the race review pipeline, run on fictional sample data.

    python3 test_race_review.py

Needs python-pptx + Pillow. Never touches the network (slides come from the
bundled library).
"""
import json
import os
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import analyze  # noqa: E402
import make_sample_data  # noqa: E402

FAIL = []


def check(label, got, want):
    if got != want:
        FAIL.append("{}\n    got:  {!r}\n    want: {!r}".format(label, got, want))


def run(*args):
    return subprocess.run([sys.executable] + list(args), capture_output=True, text=True)


tmp = tempfile.mkdtemp()
paths = make_sample_data.build(os.path.join(tmp, "sample"))

# ---- column detection --------------------------------------------------------
m = analyze.detect_columns(["First Name", "Last Name", "Email", "Gender", "Age", "Zip",
                            "Category Entered / Merchandise Ordered", "Price", "Registration Date"], {})
check("promoter export: category by prefix", m.get("category"), "Category Entered / Merchandise Ordered")
check("promoter export: reg date", m.get("reg_date"), "Registration Date")
check("promoter export: email never mapped", "Email" in m.values(), False)
m = analyze.detect_columns(["Event", "Event ID", "First Name", "Last Name", "Race Category", "Gender",
                            "Team", "City", "State", "Reg Date"], {})
check("bikereg-registrations schema", (m["category"], m["reg_date"], m["event_id"]),
      ("Race Category", "Reg Date", "Event ID"))

# ---- parsing helpers ---------------------------------------------------------
check("US datetime", str(analyze.parse_date("05/23/2025 03:14 PM")), "2025-05-23")
check("ISO date", str(analyze.parse_date("2025-05-23")), "2025-05-23")
check("money", analyze.parse_money("$1,055.00"), 1055.0)
check("gender from category", [analyze.gender_from_category(c) for c in
                               ("Women Cat 4/5", "Men Pro/1/2", "Kids 10 & Under", "Masters Women 40+")],
      ["F", "M", "U", "F"])
check("family", [analyze.category_family(c) for c in
                 ("Masters Men 50+", "Junior 15-18", "Men Pro/1/2", "Men Cat 4", "Men Beginner", "Kids 10 & Under")],
      ["Masters", "Juniors", "Pro / Elite", "Amateur (Cat 2–4)", "Beginner / Open", "Kids / Youth"])
check("merch excluded", bool(analyze.NON_RACE.search("Event T-Shirt")), True)
check("race not excluded", bool(analyze.NON_RACE.search("Men Cat 4")), False)

# ---- full analysis -----------------------------------------------------------
out = os.path.join(tmp, "analysis.json")
p = run(os.path.join(HERE, "analyze.py"), paths[2023], paths[2024], paths[2025], "-o", out)
check("analyze exit", p.returncode, 0)
a = json.load(open(out))
check("editions in order", [e["label"] for e in a["editions"]], ["2023", "2024", "2025"])
check("current/previous", (a["current"], a["previous"]), ("2025", "2024"))
want = {y: sum(make_sample_data.CATS[y].values()) for y in (2023, 2024, 2025)}
check("entries per edition", {int(k): v["entries"] for k, v in a["kpis"].items()}, want)
check("t-shirt rows excluded", a["data_quality"]["excluded_rows"], {"Event T-Shirt": 23})
check("all optional data found", all(a["flags"][k] for k in
                                     ("has_age", "has_price", "has_retention", "has_reg_date", "has_location")), True)
check("price only in one year", a["flags"]["price_yoy"], False)
check("category rename flagged", a["tables"]["category_changes"],
      {"new": ["Men Beginner", "Women Beginner"], "dropped": ["Men Cat 5"]})
check("small field found", a["tables"]["small_fields"]["rows"], [["Junior 15-18", 4, "Juniors"]])
check("price-change spike found", any(o["type"] == "reg_spike" for o in a["outliers"]), True)
check("Masters 50+ swing found", any(o["type"] == "category_swing" and "Masters Men 50+" in o["text"]
                                     for o in a["outliers"]), True)
r = a["tables"]["retention"]
check("retention adds up", r["returning"] + r["new"], a["kpis"]["2025"]["riders"])
check("home state", a["kpis"]["2025"]["home_state"], "TX")
blob = open(out).read()
check("no emails in output", "@example.com" in blob, False)
check("no rider surnames in output", any(n in blob for n in ("Lindqvist", "Okafor", "Kowalski")), False)

# single file, single year: no YoY slides, no crash
out1 = os.path.join(tmp, "one.json")
p = run(os.path.join(HERE, "analyze.py"), paths[2024], "-o", out1)
check("single-year exit", p.returncode, 0)
one = json.load(open(out1))
check("single-year has no previous", one["previous"], None)
check("single-year no movers", "category_movers" in one["charts"], False)

# explicit editions from one combined file
comb = os.path.join(tmp, "combined.csv")
with open(comb, "w") as f:
    f.write(open(paths[2023]).read())
    f.write("".join(open(paths[2024]).readlines()[1:]))
out2 = os.path.join(tmp, "comb.json")
p = run(os.path.join(HERE, "analyze.py"), comb, "-o", out2, "--edition", "Year A=61234", "--edition", "Year B=67890")
check("--edition exit", p.returncode, 0)
check("--edition labels", [e["label"] for e in json.load(open(out2))["editions"]], ["Year A", "Year B"])

p = run(os.path.join(HERE, "analyze.py"), comb, "-o", out2, "--edition", "Year A=61234")
check("unassigned rows rejected", p.returncode != 0 and "matches no --edition" in p.stderr, True)

# ---- palette ----------------------------------------------------------------
try:
    import palette
    pal_out = os.path.join(tmp, "palette.json")
    p = run(os.path.join(HERE, "palette.py"), os.path.join(tmp, "sample", "logo.png"), "-o", pal_out)
    check("palette exit", p.returncode, 0)
    pal = json.load(open(pal_out))
    check("primary is the logo orange", pal["primary"], "#F26522")
    check("secondary is the logo navy", pal["secondary"], "#12345F")
    t = palette.for_background(pal, "#131419")
    bg = palette.rgb_of("#131419")
    check("chart colors readable on dark bg",
          all(palette.contrast(palette.rgb_of(c), bg) >= 3.0 for c in t["chart"]), True)
    pw = run(os.path.join(HERE, "palette.py"), os.path.join(tmp, "sample", "logo_white.jpg"),
             "-o", os.path.join(tmp, "pw.json"))
    got = palette.rgb_of(json.load(open(os.path.join(tmp, "pw.json")))["primary"])
    check("white-background logo: background ignored (JPEG ~ orange)",
          palette.dist(got, palette.rgb_of("#F26522")) < 20, True)
    mono = palette.build_palette([((0, 0, 0), 0.6), ((255, 255, 255), 0.4)], {})
    check("monochrome logo flagged", mono["monochrome"], True)
except ImportError:
    FAIL.append("Pillow not installed — palette tests skipped")

# ---- slide library + deck ---------------------------------------------------
lib = json.load(open(os.path.join(os.path.dirname(HERE), "slides", "deck.json")))
try:
    import build_deck
    check("bundled library valid", build_deck.validate(lib), [])
    for s in lib["slides"]:
        if s.get("chart") and s["chart"] not in a["charts"]:
            FAIL.append(f"slide {s['id']} references chart {s['chart']!r} that analyze.py never produces")
        if s.get("table") and s["table"] not in a["tables"]:
            FAIL.append(f"slide {s['id']} references table {s['table']!r} that analyze.py never produces")
    for pool in lib.get("backgrounds", {}).values():
        for name in pool:
            check(f"background {name} bundled",
                  os.path.exists(os.path.join(os.path.dirname(HERE), "slides", "backgrounds", name)), True)

    deck = os.path.join(tmp, "deck.pptx")
    p = run(os.path.join(HERE, "build_deck.py"), "--analysis", out, "--palette", pal_out,
            "--logo", os.path.join(tmp, "sample", "logo.png"), "--slides", "bundled", "-o", deck)
    check("build exit", p.returncode, 0)
    if p.returncode:
        FAIL.append(p.stderr[-2000:])
    from pptx import Presentation
    prs = Presentation(deck)
    n_charts = sum(1 for sl in prs.slides for sh in sl.shapes if sh.has_chart)
    check("deck has many slides", len(prs.slides) >= 25, True)
    check("deck has native charts", n_charts >= 12, True)
    check("revenue-by-year skipped (one year of prices)", "revenue (price_yoy)" in p.stdout, True)

    deck1 = os.path.join(tmp, "one.pptx")
    p = run(os.path.join(HERE, "build_deck.py"), "--analysis", out1, "--slides", "bundled", "-o", deck1)
    check("single-year build exit (no palette, no logo)", p.returncode, 0)
    check("single-year skips YoY", "entries_trend (multi_edition)" in p.stdout, True)

    plan = run(os.path.join(HERE, "build_deck.py"), "--analysis", out, "--slides", "bundled", "--plan")
    check("--plan is JSON", isinstance(json.loads(plan.stdout)["slides"], list), True)
    check("nice_step", [build_deck.nice_step(x) for x in (4.2, 0.043, 63, 0.18)], [5, 0.05, 100, 0.2])
except ImportError:
    FAIL.append("python-pptx not installed — deck tests skipped")

if FAIL:
    print("FAILED ({}):\n".format(len(FAIL)) + "\n\n".join(FAIL))
    sys.exit(1)
print("all offline tests passed")
