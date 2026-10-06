#!/usr/bin/env python3
"""Build the race review .pptx from analysis + palette + logo + slide library.

    python3 build_deck.py --analysis analysis.json --palette palette.json \
        --logo logo.png --narrative narrative.json -o "Race Review 2025.pptx"

    python3 build_deck.py --analysis analysis.json --plan     # which slides will appear

The slide library (slides/deck.json + slides/backgrounds/) is fetched from the
plugin's GitHub repo at run time so edits there reach every user without a
plugin update. If GitHub can't be reached, or the fetched library needs a newer
builder than this one, the copy bundled next to this script is used instead.
--slides takes "auto" (default), "bundled", a local folder or deck.json, a raw
URL to a deck.json, or "owner/repo[@ref][:path/to/slides]" for a fork.

Requires python-pptx (which brings Pillow).
"""

import argparse
import json
import math
import os
import re
import ssl
import sys
import tempfile
import urllib.error
import urllib.request
from datetime import date

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt, Emu
    from PIL import Image
except ImportError:
    sys.exit("python-pptx is required: pip install python-pptx")

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import palette as palmod  # noqa: E402

BUILDER_SCHEMA = 1
LAYOUTS = {"title", "section", "bullets", "kpis", "chart", "table", "closing"}
REPO_SLIDES = "drivewayseries/bikereg-tools@main:plugins/race-review-deck/skills/race-review-deck/slides"
BUNDLED = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "slides")

W, H = 13.333, 7.5
MARGIN = 0.6


# --- slide library -----------------------------------------------------------

def raw_base(spec):
    """owner/repo[@ref][:path] -> raw.githubusercontent.com base URL ending in /"""
    m = re.match(r"^([\w.-]+)/([\w.-]+)(?:@([\w./-]+?))?(?::(.+))?$", spec)
    if not m:
        return None
    owner, repo, ref, path = m.groups()
    return f"https://raw.githubusercontent.com/{owner}/{repo}/{ref or 'main'}/{(path or 'slides').strip('/')}/"


def http_get(url, timeout=15):
    ctx = None
    try:
        import certifi
        ctx = ssl.create_default_context(cafile=certifi.where())
    except ImportError:
        pass
    req = urllib.request.Request(url, headers={"User-Agent": "race-review-deck"})
    with urllib.request.urlopen(req, timeout=timeout, context=ctx) as r:
        return r.read()


def validate(manifest):
    problems = []
    if not isinstance(manifest, dict) or not isinstance(manifest.get("slides"), list):
        return ["not a slide library (no 'slides' list)"]
    if int(manifest.get("schema", 1)) > BUILDER_SCHEMA:
        problems.append(f"needs builder schema {manifest['schema']}, this builder is {BUILDER_SCHEMA}")
    ids = set()
    for s in manifest["slides"]:
        if s.get("layout") not in LAYOUTS:
            problems.append(f"slide {s.get('id')!r}: unknown layout {s.get('layout')!r}")
        if s.get("id") in ids:
            problems.append(f"duplicate slide id {s.get('id')!r}")
        ids.add(s.get("id"))
    return problems


class Library:
    def __init__(self, manifest, base_url, local_dir, source):
        self.manifest, self.base_url, self.local_dir, self.source = manifest, base_url, local_dir, source
        self.cache = tempfile.mkdtemp(prefix="race-review-bg-")
        self.missing = []

    def background(self, name):
        """Local path for a background image: remote first (when the library is
        remote), then the bundled copy."""
        if not name:
            return None
        if self.base_url:
            dest = os.path.join(self.cache, os.path.basename(name))
            if os.path.exists(dest):
                return dest
            try:
                data = http_get(self.base_url + "backgrounds/" + name)
                with open(dest, "wb") as f:
                    f.write(data)
                return dest
            except Exception:
                pass
        for d in (self.local_dir, BUNDLED):
            if d:
                p = os.path.join(d, "backgrounds", name)
                if os.path.exists(p):
                    return p
        self.missing.append(name)
        return None


def load_library(spec):
    def local(path, label):
        if os.path.isdir(path):
            path = os.path.join(path, "deck.json")
        with open(path) as f:
            m = json.load(f)
        probs = validate(m)
        if probs:
            sys.exit(f"Slide library {path} is invalid: " + "; ".join(probs))
        return Library(m, None, os.path.dirname(os.path.abspath(path)), label)

    if spec == "bundled":
        return local(BUNDLED, "bundled with the plugin")
    if spec not in ("auto",) and not spec.startswith("http") and os.path.exists(spec):
        return local(spec, f"local: {spec}")

    if spec == "auto":
        base = raw_base(REPO_SLIDES)
    elif spec.startswith("http"):
        base = spec.rsplit("/", 1)[0] + "/" if spec.endswith(".json") else spec.rstrip("/") + "/"
    else:
        base = raw_base(spec)
        if not base:
            sys.exit(f"--slides {spec!r}: not a path, URL, or owner/repo[@ref][:path]")
    url = base + "deck.json"
    try:
        m = json.loads(http_get(url))
        probs = validate(m)
        if probs:
            print(f"note: fetched slide library at {url} can't be used ({'; '.join(probs)}); "
                  "using the bundled copy. Update the plugin to get the newer builder.")
            return local(BUNDLED, "bundled (fetched copy needs a newer builder)")
        return Library(m, base, None, f"GitHub: {url}")
    except (urllib.error.URLError, OSError, ValueError) as e:
        if spec != "auto":
            sys.exit(f"Couldn't fetch slide library {url}: {e}")
        why = f"HTTP {e.code} — not published at that path yet?" if isinstance(e, urllib.error.HTTPError) \
            else getattr(e, "reason", None) or e.__class__.__name__
        print(f"note: couldn't fetch the latest slides from GitHub ({why}); using the bundled copy.")
        return local(BUNDLED, "bundled (GitHub unreachable)")


# --- data access -------------------------------------------------------------

def dig(obj, path):
    for part in path.split("."):
        if isinstance(obj, dict) and part in obj:
            obj = obj[part]
        else:
            return None
    return obj


def nonempty(v):
    if v is None:
        return False
    if isinstance(v, dict) and "series" in v:
        return any(any(x not in (None, 0) for x in s["values"]) for s in v["series"]) and bool(v["categories"])
    if isinstance(v, (list, dict, str)):
        return len(v) > 0
    return bool(v)


class Ctx:
    def __init__(self, analysis, narrative, pal_theme):
        self.a, self.n, self.t = analysis, narrative or {}, pal_theme
        self.cur, self.prev = analysis["current"], analysis.get("previous")
        self.editions = [e["label"] for e in analysis["editions"]]
        names = sorted({n for e in analysis["editions"] for n in e["event_names"]})
        guess = re.sub(r"\s*\b(19|20)\d\d\b\s*", " ", names[-1]).strip(" -–·") if names else "Race"
        self.vars = {
            "race_name": self.n.get("race_name") or guess,
            "current": self.cur,
            "previous": self.prev or "last year",
            "editions": " · ".join(self.editions),
        }

    def fmt(self, s):
        if not s:
            return s
        try:
            return s.format(**self.vars)
        except (KeyError, IndexError, ValueError):
            return s

    def get(self, path):
        if path == "facts":
            return self.a.get("facts")
        if path == "outliers":
            return [o["text"] for o in self.a.get("outliers", [])]
        if path == "data_notes":
            return data_notes(self.a)
        if path.startswith("narrative."):
            return dig(self.n, path[len("narrative."):])
        if path in self.a.get("flags", {}):
            return self.a["flags"][path]
        return dig(self.a, path)

    def requires_ok(self, slide):
        missing = [r for r in slide.get("requires", []) if not nonempty(self.get(r))]
        return not missing, missing

    def slide_narr(self, sid):
        return (self.n.get("slides") or {}).get(sid) or {}


def data_notes(a):
    out = []
    for e in a["editions"]:
        out.append(f"{e['label']}: {', '.join(e['event_names'] or e['files'])} — "
                   f"{a['kpis'][e['label']]['entries']:,} entries"
                   + (f", event date {e['event_date']}" + (" (estimated from last registration)"
                                                          if e['event_date_inferred'] else "")
                      if e.get("event_date") else ""))
    out.append("Counts are entries, not riders: someone racing two categories counts twice. "
               "Unique-rider figures match riders by first + last name.")
    src = set(a["data_quality"]["gender_source"].values())
    if "category" in src:
        out.append("Gender is inferred from category names (Women/Men); open and kids fields are 'Unspecified'.")
    ex = a["data_quality"].get("excluded_rows") or {}
    if ex:
        out.append(f"Excluded {sum(ex.values())} non-race rows (merchandise, donations, blank categories).")
    if a["flags"].get("has_location"):
        out.append(f"Home state taken as {a['kpis'][a['current']]['home_state']} (most common rider state).")
    out.append("Source: BikeReg registration data.")
    return out


# --- drawing primitives ------------------------------------------------------

def rgb(h):
    return RGBColor.from_string(h.lstrip("#").upper())


def rect(slide, x, y, w, h, color, opacity=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(color)
    s.line.fill.background()
    s.shadow.inherit = False
    if opacity is not None and opacity < 1:
        srgb = s.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        a = srgb.makeelement(qn("a:alpha"), {"val": str(int(opacity * 100000))})
        srgb.append(a)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def text(slide, x, y, w, h, content, size=18, color="#FFFFFF", bold=False, font=None,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    lines = content if isinstance(content, list) else [content]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        r = p.add_run()
        r.text = str(line)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
        if font:
            r.font.name = font
    return tb


def bullets(slide, x, y, w, h, items, size, color, bullet_color, font, space_after=10):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.32))))
        pPr.set("indent", str(-Emu(Inches(0.32))))
        bc = pPr.makeelement(qn("a:buClr"), {})
        clr = bc.makeelement(qn("a:srgbClr"), {"val": bullet_color.lstrip("#").upper()})
        bc.append(clr)
        pPr.append(bc)
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "■"}))
        r = p.add_run()
        r.text = str(item)
        r.font.size = Pt(size)
        r.font.color.rgb = rgb(color)
        r.font.name = font
    return tb


def background(slide, lib, theme, img_name):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(theme["background_color"])
    path = lib.background(img_name)
    if path:
        pic = slide.shapes.add_picture(path, 0, 0, Inches(W), Inches(H))
        tree = slide.shapes._spTree
        tree.remove(pic._element)
        tree.insert(2, pic._element)


def place_logo(slide, logo, x, y, max_w, max_h, align="left", badge=False):
    if not logo:
        return
    lw, lh = logo["size"]
    scale = min(max_w / lw, max_h / lh)
    w, h = lw * scale, lh * scale
    if align == "center":
        x = x + (max_w - w) / 2
    elif align == "right":
        x = x + max_w - w
    if badge:
        pad = min(w, h) * 0.18
        rect(slide, x - pad, y - pad, w + 2 * pad, h + 2 * pad, "#FFFFFF", shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    slide.shapes.add_picture(logo["path"], Inches(x), Inches(y), Inches(w), Inches(h))


def notes(slide, s):
    if s:
        slide.notes_slide.notes_text_frame.text = s


# --- charts ------------------------------------------------------------------

KIND = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "bar_diverging": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
    "stacked_100": XL_CHART_TYPE.COLUMN_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


# Series that are "everything else" rather than a finding get grey, not a brand color.
NEUTRAL_SERIES = {"Unspecified", "All others", "Other", "Unknown"}


def series_colors(names, ctx):
    t = ctx.t
    eds = ctx.editions
    if names and all(n in eds for n in names):
        out = []
        for n in names:
            back = len(eds) - 1 - eds.index(n)
            out.append(t["highlight"] if back == 0 else t["prior"] if back == 1 else t["prior_2"])
        return out
    pal = t["chart"]
    out, k = [], 0
    for n in names:
        if str(n) in NEUTRAL_SERIES:
            out.append(t["prior"])
        else:
            out.append(pal[k % len(pal)])
            k += 1
    return out


def label_color_on(fill_hex, t):
    fill = palmod.rgb_of(fill_hex)
    return "#FFFFFF" if palmod.contrast((255, 255, 255), fill) >= palmod.contrast((19, 20, 25), fill) else "#131419"


def add_chart(slide, spec, x, y, w, h, ctx, font, max_items=None):
    t = ctx.t
    kind = spec["kind"]
    cats = list(spec["categories"])
    series = [dict(s) for s in spec["series"]]
    fmt = spec.get("number_format", "#,##0")
    limit = max_items or spec.get("top")
    if limit and len(cats) > limit and kind in ("bar", "column", "stacked_bar"):
        cats = cats[:limit]
        for s in series:
            s["values"] = s["values"][:limit]
    if kind == "stacked_100":
        totals = [sum((s["values"][i] or 0) for s in series) for i in range(len(cats))]
        for s in series:
            s["values"] = [round((v or 0) / tot, 4) if tot else None for v, tot in zip(s["values"], totals)]
        fmt = "0%"
    if kind in ("bar", "stacked_bar"):
        cats = cats[::-1]
        for s in series:
            s["values"] = s["values"][::-1]

    cd = CategoryChartData(number_format=fmt)
    cd.categories = [str(c) for c in cats]
    for s in series:
        cd.add_series(str(s["name"]), s["values"])
    gf = slide.shapes.add_chart(KIND[kind], Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.size = Pt(12)
    ch.font.name = font
    ch.font.color.rgb = rgb(t["muted"])
    names = [s["name"] for s in series]
    colors = series_colors(names, ctx)
    plot = ch.plots[0]

    multi = len(series) > 1
    ch.has_legend = multi or kind == "doughnut"
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP if kind != "doughnut" else XL_LEGEND_POSITION.RIGHT
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(13)
        ch.legend.font.color.rgb = rgb(t["text"])

    if kind == "doughnut":
        pts = plot.series[0].points
        dcol = [t["highlight"], t["prior"], t["prior_2"]] + t["chart"][1:]
        for i in range(len(cats)):
            pts[i].format.fill.solid()
            pts[i].format.fill.fore_color.rgb = rgb(dcol[i % len(dcol)])
            pts[i].format.line.color.rgb = rgb(t["background_color"])
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = fmt
        dl.number_format_is_linked = False
        dl.font.size = Pt(16)
        dl.font.bold = True
        dl.font.color.rgb = rgb("#131419")
        hole = plot._element.find(qn("c:holeSize"))
        if hole is not None:
            hole.set("val", "58")
        return gf

    for i, s in enumerate(plot.series):
        c = colors[i]
        if kind == "line":
            s.smooth = False
            s.format.line.color.rgb = rgb(c)
            s.format.line.width = Pt(3 if c == t["highlight"] else 2.25)
            s.marker.style = XL_MARKER_STYLE.CIRCLE
            s.marker.size = 7
            s.marker.format.fill.solid()
            s.marker.format.fill.fore_color.rgb = rgb(c)
            s.marker.format.line.color.rgb = rgb(c)
        else:
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = rgb(c)
            s.invert_if_negative = False
            if kind == "bar_diverging":
                for j, v in enumerate(series[i]["values"]):
                    p = s.points[j]
                    p.format.fill.solid()
                    p.format.fill.fore_color.rgb = rgb(t["highlight"] if (v or 0) >= 0 else t["negative"])

    if kind not in ("line",):
        plot.gap_width = 40 if len(cats) > 30 else 70
        if kind in ("column", "bar", "bar_diverging") and multi:
            plot.overlap = -10
        if kind.startswith("stacked"):
            plot.overlap = 100

    # Axes
    va, ca = ch.value_axis, ch.category_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(t["gridline"])
    va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(11)
    va.tick_labels.font.color.rgb = rgb(t["muted"])
    va.tick_labels.number_format = fmt
    va.tick_labels.number_format_is_linked = False
    ca.format.line.color.rgb = rgb(t["gridline"])
    ca.tick_labels.font.size = Pt(12 if len(cats) <= 12 else 10)
    ca.tick_labels.font.color.rgb = rgb(t["text"])
    ca.has_major_gridlines = False
    if kind == "stacked_100":
        va.maximum_scale = 1.0
        va.minimum_scale = 0
        va.major_unit = 0.25
    else:
        if kind.startswith("stacked"):
            vals = [sum((s["values"][i] or 0) for s in series) for i in range(len(cats))]
        else:
            vals = [v for s in series for v in s["values"] if v is not None]
        lo, hi = min(vals + [0]), max(vals + [0])
        if hi > lo:
            step = nice_step((hi - lo) / 5)
            # Round away float noise (0.25000000000000006) — Keynote drops tick labels on it.
            if step >= 1:
                va.major_unit = round(step, 10)
                va.maximum_scale = round(math.ceil(round(hi / step, 9)) * step, 10) if hi > 0 else 0
                va.minimum_scale = round(math.floor(round(lo / step, 9)) * step, 10) if lo < 0 else 0
            elif lo >= 0:
                # Fractions (percent axes): Keynote ignores a fractional majorUnit and splits
                # the axis into quarters, so pick a max whose quarters are whole percents.
                va.maximum_scale = next((m for m in PCT_MAXES if m >= hi - 1e-9), round(hi, 4))
    if spec.get("x_title"):
        ca.has_title = True
        ca.axis_title.text_frame.text = spec["x_title"]
        r = ca.axis_title.text_frame.paragraphs[0].runs[0]
        r.font.size = Pt(11)
        r.font.bold = False
        r.font.color.rgb = rgb(t["muted"])
    if len(cats) > 20:
        skip = max(1, len(cats) // 12)
        cat_ax = ca._element
        nm = cat_ax.find(qn("c:noMultiLvlLbl"))
        for tag in ("c:tickLblSkip", "c:tickMarkSkip"):
            el = cat_ax.makeelement(qn(tag), {"val": str(skip)})
            if nm is not None:
                nm.addprevious(el)
            else:
                cat_ax.append(el)

    # Data labels where they help and won't collide
    few = len(cats) <= 16 and len(series) <= 3
    if kind == "line":
        show = len(cats) <= 8
    elif kind in ("stacked_100", "stacked_column", "stacked_bar"):
        show = len(cats) <= 12
    else:
        show = few
    if show:
        for i, s in enumerate(plot.series):
            s_dl = s.data_labels
            s_dl.number_format = fmt
            s_dl.number_format_is_linked = False
            s_dl.font.size = Pt(11 if len(cats) > 8 else 12)
            if kind.startswith("stacked"):
                s_dl.position = XL_LABEL_POSITION.CENTER
                s_dl.font.color.rgb = rgb(label_color_on(colors[i], t))
                vals = series[i]["values"]
                tot_scale = 1 if kind == "stacked_100" else max(
                    (sum((ss["values"][j] or 0) for ss in series) for j in range(len(cats))), default=1) or 1
                for j, v in enumerate(vals):
                    if not v or (v / tot_scale) < 0.05:
                        s.points[j].data_label.has_text_frame = True
                        s.points[j].data_label.text_frame.text = ""
            elif kind == "line":
                s_dl.position = XL_LABEL_POSITION.ABOVE
                s_dl.font.color.rgb = rgb(t["text"])
            else:
                s_dl.position = XL_LABEL_POSITION.OUTSIDE_END
                s_dl.font.color.rgb = rgb(t["text"])
            s_dl.show_value = True
    return gf


PCT_MAXES = [0.04, 0.08, 0.12, 0.16, 0.2, 0.24, 0.28, 0.32, 0.4, 0.6, 0.8, 1.0, 1.2, 1.6, 2.0]


def nice_step(raw):
    """Round an axis step to 1, 2, 2.5 or 5 x 10^k so ticks land on readable values."""
    if raw <= 0:
        return 1
    mag = 10 ** math.floor(math.log10(raw))
    for m in (1, 2, 2.5, 5, 10):
        if raw <= m * mag:
            return m * mag
    return 10 * mag


# --- tables ------------------------------------------------------------------

def fmt_cell(v, header):
    if v is None:
        return "—"
    h = str(header).lower()
    if isinstance(v, float) and ("%" in h or "share" in h):
        return f"{'+' if v > 0 and '%' in h else ''}{round(v * 100)}%"
    if h == "change" and isinstance(v, (int, float)):
        return f"{'+' if v > 0 else ''}{v:,.0f}"
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, int):
        return f"{v:,}"
    if isinstance(v, float):
        return f"{v:,.1f}"
    return str(v)


def add_table(slide, data, x, y, w, ctx, font, max_rows=12):
    t = ctx.t
    cols = data["columns"]
    rows = data["rows"][:max_rows]
    rh = 0.4
    gt = slide.shapes.add_table(len(rows) + 1, len(cols), Inches(x), Inches(y), Inches(w),
                                Inches(rh * (len(rows) + 1)))
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid" — we paint cells ourselves
    first = w * (0.42 if len(cols) > 2 else 0.6)
    rest = (w - first) / max(1, len(cols) - 1)
    for i in range(len(cols)):
        tbl.columns[i].width = Inches(first if i == 0 else rest)
    for r in range(len(rows) + 1):
        tbl.rows[r].height = Inches(rh)
        for c in range(len(cols)):
            cell = tbl.cell(r, c)
            if r == 0:
                val, fill, color, bold = str(cols[c]), t["band"], t["on_band"], True
            else:
                val = fmt_cell(rows[r - 1][c], cols[c])
                fill = t["row_alt"] if r % 2 == 0 else t["background_color"]
                color, bold = t["text"], False
                if str(cols[c]).lower() == "change" or "%" in str(cols[c]):
                    raw = rows[r - 1][c]
                    if isinstance(raw, (int, float)) and raw:
                        color = t["positive"] if raw > 0 else t["negative"]
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.text = val
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
            run = p.runs[0]
            run.font.size = Pt(13)
            run.font.bold = bold
            run.font.name = font
            run.font.color.rgb = rgb(color)
            # Hairline borders in the gridline color
            tcPr = cell._tc.get_or_add_tcPr()
            for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = tcPr.makeelement(qn(side), {"w": "6350"})
                sf = ln.makeelement(qn("a:solidFill"), {})
                sf.append(sf.makeelement(qn("a:srgbClr"), {"val": t["gridline"].lstrip("#")}))
                ln.append(sf)
                tcPr.append(ln)
    extra = len(data["rows"]) - len(rows)
    return gt, extra


# --- slide renderers ---------------------------------------------------------

class Deck:
    def __init__(self, lib, ctx, logo, theme):
        self.lib, self.ctx, self.logo, self.theme = lib, ctx, logo, theme
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.rot = {}
        self.section_no = 0
        self.page = 0
        self.hf = theme.get("heading_font", "Arial")
        self.bf = theme.get("body_font", "Arial")

    def bg_for(self, layout, s):
        if s.get("background"):
            return s["background"]
        pool = (self.lib.manifest.get("backgrounds") or {}).get(layout) or \
               (self.lib.manifest.get("backgrounds") or {}).get("content") or []
        if not pool:
            return None
        i = self.rot.get(layout, 0)
        self.rot[layout] = i + 1
        return pool[i % len(pool)]

    def new(self, layout, s):
        slide = self.prs.slides.add_slide(self.blank)
        background(slide, self.lib, self.theme, self.bg_for(layout, s))
        self.page += 1
        return slide

    def chrome(self, slide, headline, extra_note=None):
        t = self.ctx.t
        rect(slide, MARGIN, 0.42, 0.55, 0.07, t["highlight"])
        text(slide, MARGIN, 0.58, W - 2 * MARGIN - 1.6, 1.05, headline, size=26, bold=True,
             color=t["text"], font=self.hf, anchor=MSO_ANCHOR.TOP)
        foot = self.ctx.fmt(self.lib.manifest.get("footer", ""))
        if extra_note:
            foot = f"{foot}  ·  {extra_note}" if foot else extra_note
        text(slide, MARGIN, H - 0.45, W - 2 * MARGIN - 1.0, 0.3, foot, size=10, color=t["muted"], font=self.bf)
        text(slide, W - MARGIN - 0.6, H - 0.45, 0.6, 0.3, str(self.page), size=10, color=t["muted"],
             font=self.bf, align=PP_ALIGN.RIGHT)
        if self.logo:
            place_logo(slide, self.logo, W - MARGIN - 1.3, 0.42, 1.3, 0.55, align="right",
                       badge=self.logo["badge"])

    def panel(self, slide, x, y, w, h):
        op = float(self.theme.get("content_panel_opacity", 0.82))
        rect(slide, x, y, w, h, self.theme["background_color"], opacity=op,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)

    def headline(self, s):
        n = self.ctx.slide_narr(s["id"])
        return self.ctx.fmt(n.get("headline") or s.get("title", ""))

    # layouts
    def title(self, s):
        t = self.ctx.t
        slide = self.new("title", s)
        n = self.ctx.slide_narr(s["id"])
        if self.logo:
            place_logo(slide, self.logo, 0.9, 0.9, 4.2, 2.0, badge=self.logo["badge"])
        rect(slide, 0.9, 3.55, 0.9, 0.09, t["highlight"])
        text(slide, 0.9, 3.8, 8.5, 1.4, self.ctx.fmt(n.get("headline") or s.get("title")),
             size=50, bold=True, color=t["text"], font=self.hf, anchor=MSO_ANCHOR.TOP)
        text(slide, 0.9, 5.15, 8.5, 0.6, self.ctx.fmt(n.get("subtitle") or self.ctx.n.get("subtitle")
                                                       or s.get("subtitle", "")),
             size=26, color=t["highlight"], font=self.hf)
        line = self.ctx.n.get("date_line")
        if not line:
            ed = next(e for e in self.ctx.a["editions"] if e["label"] == self.ctx.cur)
            if ed.get("event_date") and not ed.get("event_date_inferred"):
                line = date.fromisoformat(ed["event_date"]).strftime("%B %-d, %Y")
            elif len(self.ctx.editions) > 1:
                line = "Compared with " + ", ".join(self.ctx.editions[:-1])
        if line:
            text(slide, 0.9, 5.8, 8.5, 0.5, line, size=16, color=t["muted"], font=self.bf)
        notes(slide, n.get("notes"))

    def section(self, s):
        t = self.ctx.t
        slide = self.new("section", s)
        self.section_no += 1
        n = self.ctx.slide_narr(s["id"])
        text(slide, 0.9, 2.55, 3, 0.5, f"{self.section_no:02d}", size=20, bold=True, color=t["highlight"],
             font=self.hf)
        text(slide, 0.9, 3.0, 9, 1.2, self.ctx.fmt(n.get("headline") or s["title"]), size=48, bold=True,
             color=t["text"], font=self.hf)
        sub = n.get("subtitle") or s.get("subtitle")
        if sub:
            text(slide, 0.9, 4.2, 9, 0.8, self.ctx.fmt(sub), size=20, color=t["muted"], font=self.bf)
        notes(slide, n.get("notes"))

    def closing(self, s):
        t = self.ctx.t
        slide = self.new("closing", s)
        n = self.ctx.slide_narr(s["id"])
        if self.logo:
            place_logo(slide, self.logo, W / 2 - 2.0, 1.5, 4.0, 2.0, align="center", badge=self.logo["badge"])
        text(slide, 1, 4.0, W - 2, 1.0, self.ctx.fmt(n.get("headline") or s.get("title", "Thank you")),
             size=44, bold=True, color=t["text"], font=self.hf, align=PP_ALIGN.CENTER)
        sub = n.get("subtitle") or s.get("subtitle")
        if sub:
            text(slide, 1, 5.0, W - 2, 0.6, self.ctx.fmt(sub), size=20, color=t["highlight"], font=self.bf,
                 align=PP_ALIGN.CENTER)
        notes(slide, n.get("notes"))

    def bullets_slide(self, s, items=None):
        t = self.ctx.t
        n = self.ctx.slide_narr(s["id"])
        if items is None:
            items = n.get("bullets")
        if not items:
            for src in s.get("items", []):
                v = self.ctx.get(src)
                if nonempty(v):
                    items = v
                    break
        items = [i["text"] if isinstance(i, dict) else i for i in (items or [])][: s.get("max_items", 8)]
        if not items:
            return False
        slide = self.new("content", s)
        self.panel(slide, MARGIN - 0.2, 1.75, W - 2 * MARGIN + 0.4, H - 2.45)
        self.chrome(slide, self.headline(s))
        small = s.get("small") or len(items) > 6 or sum(len(i) for i in items) > 650
        bullets(slide, MARGIN + 0.2, 2.05, W - 2 * MARGIN - 0.4, H - 2.9, items,
                size=15 if small else 20, color=t["text"], bullet_color=t["highlight"], font=self.bf,
                space_after=8 if small else 14)
        notes(slide, n.get("notes"))
        return True

    def kpis(self, s):
        t = self.ctx.t
        n = self.ctx.slide_narr(s["id"])
        src = s.get("source")
        vals = self.ctx.get(src) if src else self.ctx.a["kpis"][self.ctx.cur]
        deltas = {} if src else self.ctx.a.get("kpi_delta", {})
        metrics = [m for m in s["metrics"] if (vals or {}).get(m["key"]) is not None][:6]
        if not metrics:
            return False
        slide = self.new("content", s)
        self.chrome(slide, self.headline(s))
        cols = 3 if len(metrics) in (3, 5, 6) else 2 if len(metrics) in (2, 4) else 1
        rows = (len(metrics) + cols - 1) // cols
        side = 3.9 if n.get("bullets") else 0
        gx, gy, gw, gh = MARGIN, 1.85, W - 2 * MARGIN - side, H - 2.55
        gap = 0.25
        tw = (gw - gap * (cols - 1)) / cols
        th = min(2.3, (gh - gap * (rows - 1)) / rows)
        for i, m in enumerate(metrics):
            r, c = divmod(i, cols)
            x, y = gx + c * (tw + gap), gy + r * (th + gap)
            self.panel(slide, x, y, tw, th)
            rect(slide, x, y + 0.25, 0.07, th - 0.5, t["highlight"])
            v = vals[m["key"]]
            text(slide, x + 0.35, y + 0.22, tw - 0.5, 0.95, fmt_metric(v, m.get("format")), size=40, bold=True,
                 color=t["text"], font=self.hf)
            text(slide, x + 0.35, y + 1.12, tw - 0.5, 0.5, self.ctx.fmt(m["label"]), size=15, color=t["muted"],
                 font=self.bf)
            d = deltas.get(m["key"])
            if d and self.ctx.prev:
                txt, up = fmt_delta(d, m.get("format"))
                if txt:
                    text(slide, x + 0.35, y + 1.55, tw - 0.5, 0.4, f"{txt} vs {self.ctx.prev}", size=13,
                         bold=True, color=t["positive"] if up else t["negative"] if up is False else t["muted"],
                         font=self.bf)
        if side:
            self.takeaway(slide, n["bullets"], W - MARGIN - side + 0.25, gy, side - 0.25, gh)
        notes(slide, n.get("notes"))
        return True

    def takeaway(self, slide, items, x, y, w, h):
        t = self.ctx.t
        self.panel(slide, x, y, w, h)
        rect(slide, x, y, w, 0.07, t["highlight"])
        text(slide, x + 0.3, y + 0.25, w - 0.6, 0.4, "TAKEAWAYS", size=12, bold=True, color=t["highlight"],
             font=self.hf)
        bullets(slide, x + 0.3, y + 0.7, w - 0.6, h - 0.9, items[:5], size=14, color=t["text"],
                bullet_color=t["highlight"], font=self.bf, space_after=10)

    def chart(self, s):
        n = self.ctx.slide_narr(s["id"])
        spec = self.ctx.get("charts." + s["chart"])
        if not nonempty(spec):
            return False
        slide = self.new("content", s)
        self.chrome(slide, self.headline(s), extra_note=spec.get("note"))
        side = 3.9 if n.get("bullets") else 0
        cx, cy, cw, chh = MARGIN - 0.2, 1.75, W - 2 * MARGIN + 0.4 - side, H - 2.45
        self.panel(slide, cx, cy, cw, chh)
        add_chart(slide, spec, cx + 0.2, cy + 0.15, cw - 0.4, chh - 0.3, self.ctx, self.bf, s.get("max_items"))
        if side:
            self.takeaway(slide, n["bullets"], W - MARGIN - side + 0.45, cy, side - 0.25, chh)
        notes(slide, n.get("notes"))
        return True

    def table(self, s):
        n = self.ctx.slide_narr(s["id"])
        data = self.ctx.get("tables." + s["table"])
        if not data or not data.get("rows"):
            return False
        slide = self.new("content", s)
        side = 3.9 if n.get("bullets") else 0
        max_rows = min(s.get("max_rows", 12), 11)
        tw = W - 2 * MARGIN - side
        rows_shown = min(len(data["rows"]), max_rows)
        th = 0.4 * (rows_shown + 1)
        self.panel(slide, MARGIN - 0.2, 1.75, tw + 0.4, th + 0.4)
        _, extra = add_table(slide, data, MARGIN, 1.95, tw, self.ctx, self.bf, max_rows)
        self.chrome(slide, self.headline(s), extra_note=f"{extra} more rows not shown" if extra else None)
        if side:
            self.takeaway(slide, n["bullets"], W - MARGIN - side + 0.45, 1.75, side - 0.25, H - 2.45)
        notes(slide, n.get("notes"))
        return True


def fmt_metric(v, f):
    if f == "pct":
        return f"{round(v * 100)}%"
    if f == "money":
        return f"${v:,.0f}" if v < 100000 else f"${v / 1000:,.0f}k"
    if f == "decimal":
        return f"{v:,.1f}".rstrip("0").rstrip(".")
    return f"{v:,.0f}"


def fmt_delta(d, f):
    if d.get("unit") == "pts":
        v = d["abs"] * 100
        if abs(v) < 0.05:
            return "no change", None
        return f"{'▲' if v > 0 else '▼'} {abs(v):.1f} pts", v > 0
    p = d.get("pct")
    if p is None:
        a = d.get("abs")
        if not a:
            return "no change", None
        return f"{'▲' if a > 0 else '▼'} {abs(a):,.0f}", a > 0
    if abs(p) < 0.005:
        return "no change", None
    return f"{'▲' if p > 0 else '▼'} {abs(p) * 100:.0f}%", p > 0


# --- main --------------------------------------------------------------------

def prepare_logo(path, cache):
    if not path:
        return None
    im = Image.open(path)
    im.load()
    png = os.path.join(cache, "logo.png")
    has_alpha = im.mode in ("RGBA", "LA", "PA") or (im.mode == "P" and "transparency" in im.info)
    conv = im.convert("RGBA") if has_alpha else im.convert("RGB")
    conv.save(png, "PNG")
    # Badge the logo (white rounded card) when it would vanish or clash on a dark photo:
    # an opaque light rectangle, or dark artwork on transparency.
    rgba = conv.convert("RGBA")
    rgba.thumbnail((200, 200))
    px = [p for p in palmod.pixel_list(rgba) if p[3] > 128]
    lum = sum(palmod.luminance(p[:3]) for p in px) / max(1, len(px))
    border = palmod.border_color(im)
    badge = (border is not None and palmod.luminance(border) > 0.4) or (border is None and lum < 0.12)
    return {"path": png, "size": conv.size, "badge": badge}


def plan(lib, ctx):
    out = []
    for s in lib.manifest["slides"]:
        ok, missing = ctx.requires_ok(s)
        n = ctx.slide_narr(s["id"])
        if n.get("skip"):
            ok, missing = False, ["skipped in narrative"]
        out.append({"id": s["id"], "layout": s["layout"], "title": ctx.fmt(s.get("title")),
                    "data": s.get("chart") or s.get("table") or s.get("source") or s.get("items"),
                    "include": ok, "missing": missing})
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--analysis", required=True)
    ap.add_argument("--palette", help="palette.json from palette.py (default: neutral)")
    ap.add_argument("--logo")
    ap.add_argument("--narrative", help="narrative.json written for this race")
    ap.add_argument("--slides", default="auto")
    ap.add_argument("-o", "--output", default="race-review.pptx")
    ap.add_argument("--plan", action="store_true", help="print which slides will be built, as JSON, and exit")
    a = ap.parse_args(argv)

    with open(a.analysis) as f:
        analysis = json.load(f)
    narrative = {}
    if a.narrative:
        with open(a.narrative) as f:
            narrative = json.load(f)
    if a.palette:
        with open(a.palette) as f:
            pal = json.load(f)
    else:
        pal = palmod.build_palette([], {"primary": "#3B82F6"})

    lib = load_library(a.slides)
    theme = dict(lib.manifest.get("theme") or {})
    theme.setdefault("background_color", "#131419")
    t = palmod.for_background(pal, theme["background_color"])
    t["background_color"] = theme["background_color"]
    ctx = Ctx(analysis, narrative, t)

    p = plan(lib, ctx)
    if a.plan:
        print(json.dumps({"library": lib.source, "slides": p}, indent=1))
        return

    cache = tempfile.mkdtemp(prefix="race-review-")
    logo = prepare_logo(a.logo, cache)
    deck = Deck(lib, ctx, logo, theme)

    extras = {}
    for e in narrative.get("extra_slides", []) or []:
        extras.setdefault(e.get("after"), []).append(e)

    built, skipped = [], []
    for s, info in zip(lib.manifest["slides"], p):
        if info["include"]:
            fn = {"title": deck.title, "section": deck.section, "closing": deck.closing,
                  "bullets": deck.bullets_slide, "kpis": deck.kpis, "chart": deck.chart,
                  "table": deck.table}[s["layout"]]
            r = fn(s)
            if r is False:
                skipped.append((s["id"], "no data"))
            else:
                built.append(s["id"])
        else:
            skipped.append((s["id"], ", ".join(info["missing"])))
        for e in extras.get(s["id"], []):
            es = {"id": e.get("id", f"extra_{len(built)}"), "layout": "bullets", "title": e.get("title", "")}
            deck.ctx.n.setdefault("slides", {})[es["id"]] = {"headline": e.get("title"), "notes": e.get("notes")}
            if deck.bullets_slide(es, items=e.get("bullets")):
                built.append(es["id"])

    deck.prs.save(a.output)
    print(f"Slide library: {lib.source}")
    if lib.missing:
        print(f"WARNING: background images not found: {', '.join(sorted(set(lib.missing)))} (plain background used)")
    print(f"Built {len(built)} slides: {', '.join(built)}")
    if skipped:
        print("Skipped: " + "; ".join(f"{i} ({why})" for i, why in skipped))
    unused = set((narrative.get("slides") or {})) - set(built) - {e.get("id") for e in narrative.get("extra_slides", []) or []}
    if unused:
        print(f"note: narrative has text for slides that weren't built: {', '.join(sorted(unused))}")
    print(f"Wrote {a.output}")


if __name__ == "__main__":
    main()
